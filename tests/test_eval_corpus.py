#!/usr/bin/env python3
"""
The evaluation machinery, broken on purpose.

`tests/test_guards.py` proves the extractor's guards fire. This file does the same for
the thing that measures it, because a harness that miscounts is worse than no harness:
it produces a number everything else is then steered by.

Each test breaks one specific thing -- the cache key that lost a year, the abstention
that vanished into an uncounted bucket, the interval that would call noise a finding --
and asserts the harness notices. No network, no model, no API key.

Run:  python -m pytest tests/test_eval_corpus.py -v
"""

import importlib.util
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl                                              # noqa: E402

from evalkit import (PeriodRef, cache_path_for, legacy_cache_path,   # noqa: E402
                     load_truth_sheet, Cell)
import evalmetrics as M                                      # noqa: E402
from periods import parse_period, parse_filing_stem          # noqa: E402


def _script(name):
    """Load a script by path, the way the other test modules do."""
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _cell(key="ARCI Q1 2022", field="aset", status="correct", **kw):
    ticker, quarter, year = key.split(" ")
    return Cell(key=key, ticker=ticker, quarter=quarter, year=int(year),
                field=field, status=status, **kw)


def _workbook(path: Path, ticker: str, headers, values=None):
    """A minimal ground-truth sheet: D1 names the issuer, row 3 heads the periods."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["D1"] = ticker
    for index, header in enumerate(headers):
        column = index + 2
        sheet.cell(row=3, column=column).value = header
        for row in range(4, 14):
            sheet.cell(row=row, column=column).value = (values or {}).get(header, 1000.0)
    workbook.save(path)
    return path


# --- the cache key that lost a year ----------------------------------------------

def test_the_prediction_cache_key_carries_the_year(tmp_path):
    """ARCI Q1 2022 and ARCI Q1 2023 are different filings and must not share a file.

    The old key was <TICKER>_<Qn>. With one year of filings that was invisible; with
    five, the second year silently scores the first year's prediction against its own
    labels -- every field wrong, for a reason that appears nowhere in the output.
    """
    a = PeriodRef("ARCI", "Q1", 2022, tmp_path / "Q1_2022_ARCI.pdf")
    b = PeriodRef("ARCI", "Q1", 2023, tmp_path / "Q1_2023_ARCI.pdf")
    assert cache_path_for(a, tmp_path, "tag") != cache_path_for(b, tmp_path, "tag")
    assert cache_path_for(a, tmp_path, "tag").name == "ARCI_Q1_2022.json"


def test_a_legacy_cache_entry_is_only_read_when_its_date_proves_the_period(tmp_path):
    """The un-yeared cache is still honoured -- but only when it can prove itself.

    Twenty-two committed predictions use the old name, and throwing them away would cost
    the one thing a reviewer without a key can check. They are accepted on evidence: the
    payload's own period_end_date. That makes the migration unable to repeat the
    collision it exists to fix.
    """
    cache = tmp_path / "predictions" / "tag"
    cache.mkdir(parents=True)
    (cache / "ARCI_Q1.json").write_text(json.dumps({"period_end_date": "2022-03-31"}))

    matching = PeriodRef("ARCI", "Q1", 2022, tmp_path / "x.pdf")
    other_year = PeriodRef("ARCI", "Q1", 2023, tmp_path / "y.pdf")
    assert legacy_cache_path(matching, tmp_path, "tag").name == "ARCI_Q1.json"
    assert legacy_cache_path(other_year, tmp_path, "tag") is None


def test_a_predicted_nil_is_classified_rather_than_crashing():
    """The taxonomy divides by the prediction/truth ratio, and a predicted zero made
    every one of those tests raise ZeroDivisionError -- losing the whole filing to a
    crash in the code whose only job is to describe a wrong answer. Found by the first
    full-corpus pass: 223 filings produced the case four quarters never had."""
    from evalkit import classify_failure, compare
    assert classify_failure(0.0, 1_000.0, None) == "wrong_value"
    assert compare(0.0, 1_000.0, 1e-4) == ("wrong", 1.0, "wrong_value")
    # The reverse -- a figure where the filing prints nil -- was already handled.
    assert classify_failure(1_000.0, 0.0, None) == "wrong_value"


# --- periods and corpus discovery -------------------------------------------------

def test_a_period_is_matched_as_a_tuple_not_a_spelling():
    """The workbooks head 2020-2024 "TAHUNAN <year>" and 2025 "Q4 2025".

    Code that builds a label string to look for finds no annual column at all for
    whichever spelling it did not build -- five years of fourth quarters, scored as
    nothing, with no error.
    """
    assert parse_period("TAHUNAN 2024") == parse_period("Q4 2024") == ("Q4", 2024)
    assert parse_period("Q4 2025") == ("Q4", 2025)
    assert parse_period("Harga saham rupiah") is None
    assert parse_filing_stem("Q1_2022_ARCI") == ("ARCI", "Q1", 2022)
    assert parse_filing_stem("notes_2022") is None


def test_a_workbook_labelled_for_another_issuer_is_reported_not_repaired(tmp_path):
    """A stale D1 must not end a 223-filing run, and must not be silently overridden.

    The ticker guard exists because these workbooks are made by copying each other. In a
    corpus run the right response is to name the workbook and carry on: a fatal error
    costs every issuer after it, and a silent override scores one issuer's extraction
    against another's figures -- the exact failure the guard was written for.
    """
    path = _workbook(tmp_path / "ADMR.xlsx", "ITMG", ["Q1 2022"])
    before = path.stat().st_mtime

    sheet = load_truth_sheet(path, "ADMR")
    assert sheet.problem and "ITMG" in sheet.problem and "ADMR" in sheet.problem
    assert sheet.columns == {}, "an untrusted sheet contributes no labels"
    assert path.stat().st_mtime == before, "the workbook must not be written to"


def test_a_labelled_column_is_found_under_either_annual_spelling(tmp_path):
    path = _workbook(tmp_path / "DSNG.xlsx", "DSNG", ["TAHUNAN 2022", "Q4 2025"])
    sheet = load_truth_sheet(path, "DSNG")
    assert sheet.problem is None
    assert ("Q4", 2022) in sheet.columns and ("Q4", 2025) in sheet.columns
    assert sheet.columns[("Q4", 2022)]["aset"] == 1000.0


def test_a_subset_corpus_run_gets_its_own_scorecard_name(monkeypatch, tmp_path):
    """`--tickers LSIP TAPG` must not overwrite the 223-filing scorecard.

    The single-ticker path has carried a subset marker since the day a one-period run
    replaced a four-period artifact. The corpus path shipped without one, and the first
    two-issuer run overwrote the full corpus scorecard within the hour -- recoverable
    only because the prediction cache can rebuild it.
    """
    run_eval = _script("run_eval")
    names = []
    monkeypatch.setattr(run_eval, "run_corpus",
                        lambda args, tag, out, name: names.append(name) or 0)
    monkeypatch.setenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
    monkeypatch.setenv("GEMINI_INPUT_MODE", "image")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    for argv in (["run_eval", "--corpus", "--out", str(tmp_path)],
                 ["run_eval", "--corpus", "--tickers", "LSIP", "TAPG", "--out", str(tmp_path)],
                 ["run_eval", "--corpus", "--limit", "5", "--out", str(tmp_path)]):
        monkeypatch.setattr("sys.argv", argv)
        run_eval.main()

    full, tickers, sampled = names
    assert full == "corpus_scorecard_gemini_gemini-3.1-flash-lite_image.json"
    assert tickers == "corpus_scorecard_gemini_gemini-3.1-flash-lite_image_LSIPTAPG.json"
    assert sampled.endswith("_n5s0.json")
    assert len({full, tickers, sampled}) == 3


# --- the metrics ------------------------------------------------------------------

def test_an_abstention_on_an_unlabelled_cell_is_counted_rather_than_absorbed():
    """`both_absent` is the bucket that hides things.

    A field nobody labelled, which the system withheld, used to land there -- outside
    the numerator and outside the denominator, indistinguishable from a field that was
    never asked for. Coverage separates the two.
    """
    cells = [_cell(field="aset", status="correct"),
             _cell(field="total_share", status="both_absent", abstained=True),
             _cell(field="kas", status="both_absent")]
    coverage = M.coverage(cells)
    assert coverage["labelled"] == 1
    assert coverage["unlabelled_abstained"] == 1
    assert coverage["unlabelled_silent"] == 1
    assert coverage["cells"] == 3


def test_accuracy_separates_what_it_answered_from_what_it_withheld():
    cells = [_cell(field="aset", status="correct"),
             _cell(field="kas", status="wrong"),
             _cell(field="liabilitas", status="missed", abstained=True)]
    numbers = M.accuracy(cells)
    assert numbers["accuracy"] == pytest.approx(1 / 3)
    assert numbers["accuracy_on_answered"] == pytest.approx(1 / 2)


def test_abstention_precision_counts_saves_and_losses_and_refuses_to_guess():
    """Was withholding right? Only the withheld figure can say.

    A withdrawal whose figure was never recorded is unknown, and assuming it either way
    turns a missing measurement into whichever number flatters the confidence layer.
    """
    cells = [
        _cell(field="a", status="missed", abstained=True, would_have_been="wrong"),
        _cell(field="b", status="missed", abstained=True, would_have_been="correct"),
        _cell(field="c", status="missed", abstained=True, would_have_been=None),
        _cell(field="d", status="wrong"),
    ]
    result = M.abstention(cells)
    assert (result["withdrawn"], result["withdrawn_known"]) == (3, 2)
    assert (result["caught"], result["thrown_away"]) == (1, 1)
    assert result["unknown_withheld"] == 1
    assert result["precision"] == pytest.approx(0.5)
    # One caught, one shipped wrong: the layer caught half of what would have been wrong.
    assert result["recall"] == pytest.approx(0.5)


def test_calibration_measures_overconfidence_and_ignores_withheld_cells():
    """Ten cells claiming 0.9 that are right five times is an ECE of 0.4, not a pass."""
    cells = [_cell(field=f"f{i}", status="correct" if i < 5 else "wrong", confidence=0.9)
             for i in range(10)]
    cells.append(_cell(field="withheld", status="missed", abstained=True, confidence=0.3))
    result = M.calibration(cells)
    assert result["n"] == 10, "withheld cells are not evidence about answered ones"
    bucket = next(b for b in result["buckets"] if b["n"])
    assert bucket["accuracy"] == pytest.approx(0.5)
    assert bucket["gap"] == pytest.approx(-0.4)          # negative = overconfident
    assert result["ece"] == pytest.approx(0.4)


def test_a_slice_carries_its_n_so_four_cells_cannot_read_as_a_result():
    cells = [_cell(key="ARCI Q1 2022", field="aset", status="correct"),
             _cell(key="JPFA Q1 2022", field="aset", status="wrong")]
    rows = M.by_slice(cells, M.SLICES["ticker"])
    assert rows["JPFA"]["scored"] == 1 and rows["JPFA"]["accuracy"] == 0.0
    assert rows["ARCI"]["filings"] == 1


def test_cost_is_null_with_a_reason_rather_than_zero_dollars():
    """`config/pricing.json` ships empty. A zero here would read as free."""
    filings = [{"elapsed_s": 10.0,
                "usage": {"llm_calls": 4, "tokens": {"input": 1000, "output": 100},
                          "latency_ms": {"extract": 900},
                          "cost": {"amount": None, "currency": "USD",
                                   "reason": "no price configured for gemini-3.1-flash-lite"}}}]
    result = M.cost_latency(filings)
    assert result["cost"]["amount"] is None
    assert "no price configured" in result["cost"]["reason"]
    assert result["tokens_per_filing"]["input"] == 1000
    assert result["seconds_per_filing"]["p50"] == 10.0


def test_the_report_prints_the_reason_instead_of_a_zero_cost(tmp_path):
    report_eval = _script("report_eval")
    scorecard = {
        "ticker": "CORPUS", "provider": "tag", "cells": [_cell().to_dict()],
        "metrics": M.summarise([_cell()], [{"elapsed_s": 5.0, "usage": {
            "llm_calls": 3, "tokens": {"input": 10, "output": 1}, "latency_ms": {},
            "cost": {"amount": None, "currency": "USD",
                     "reason": "no price configured for gemini-3.1-flash-lite"}}}]),
    }
    markdown = report_eval.render_markdown(scorecard)
    assert "no price configured" in markdown
    assert "$0.00" not in markdown


# --- the interval -----------------------------------------------------------------

def test_an_identical_run_has_an_interval_that_contains_zero():
    cells = [_cell(key=f"ARCI Q{q} 2022", field=f"f{i}",
                   status="correct" if i % 4 else "wrong")
             for q in (1, 2, 3, 4) for i in range(10)]
    result = M.bootstrap_delta(cells, cells, iters=500)
    assert result["delta"] == 0
    assert result["lo"] <= 0 <= result["hi"]
    assert result["n_clusters"] == 4, "the filing is the resampling unit, not the cell"


def test_a_clear_improvement_excludes_zero():
    base = [_cell(key=f"ARCI Q{q} {y}", field=f"f{i}", status="wrong")
            for y in range(2018, 2026) for q in (1, 2, 3, 4) for i in range(10)]
    better = [_cell(key=c.key, field=c.field, status="correct") for c in base]
    result = M.bootstrap_delta(base, better, iters=500)
    assert result["delta"] == pytest.approx(1.0)
    assert result["lo"] > 0 and result["p_improvement"] == 1.0


def test_the_interval_is_wider_when_issuers_rather_than_filings_are_resampled():
    """Filings of one issuer share a layout, so they are not independent evidence.

    Clustering by issuer answers the harder question -- would this hold for an issuer
    the corpus has never seen -- and must not report a narrower interval than the
    easier one.
    """
    base, better = [], []
    for ticker in ("ARCI", "JPFA", "CPIN"):
        for quarter in (1, 2, 3, 4):
            for i in range(10):
                key = f"{ticker} Q{quarter} 2022"
                fixed = ticker == "ARCI"          # only one issuer improves
                base.append(_cell(key=key, field=f"f{i}", status="wrong"))
                better.append(_cell(key=key, field=f"f{i}",
                                    status="correct" if fixed else "wrong"))
    by_filing = M.bootstrap_delta(base, better, iters=500, cluster="filing")
    by_ticker = M.bootstrap_delta(base, better, iters=500, cluster="ticker")
    assert by_filing["n_clusters"] == 12 and by_ticker["n_clusters"] == 3
    assert (by_ticker["hi"] - by_ticker["lo"]) > (by_filing["hi"] - by_filing["lo"])


# --- the driver -------------------------------------------------------------------

def test_cached_only_never_calls_a_provider(monkeypatch, tmp_path):
    """A CI gate that "should not" call a model is not a guarantee. This makes it one."""
    run_eval = _script("run_eval")
    monkeypatch.setattr(run_eval, "extract_from_pdf",
                        lambda *a, **k: pytest.fail("a provider was called"))
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
    monkeypatch.setenv("GEMINI_INPUT_MODE", "image")
    monkeypatch.setattr("sys.argv", ["run_eval", "--corpus", "--cached-only",
                                     "--out", str(tmp_path)])

    assert run_eval.main() in (0, 1)
    scorecard = json.loads(next(tmp_path.glob("corpus_scorecard_*.json")).read_text())
    assert scorecard["corpus"]["uncached"], "uncached filings should be named, not extracted"


def test_one_failing_filing_does_not_abandon_the_corpus(monkeypatch, tmp_path):
    """In a 223-filing pass, one broken PDF must cost one row, not the measurement."""
    run_eval = _script("run_eval")
    refs = [PeriodRef("ARCI", "Q1", 2022, ROOT / "data" / "raw" / "Q1_2022_ARCI.pdf"),
            PeriodRef("ARCI", "Q2", 2022, ROOT / "data" / "raw" / "Q2_2022_ARCI.pdf")]
    truth = {field: 1.0 for field in __import__("schema").EXCEL_ROWS}

    class _Sheet:
        columns = {("Q1", 2022): truth, ("Q2", 2022): truth}
    monkeypatch.setattr(run_eval, "corpus_pairs",
                        lambda *a, **k: (refs, {"ARCI": _Sheet()}, []))

    calls = {"n": 0}

    def flaky(*args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("no /Root object")
        raise run_eval.LLMError("provider is down")
    monkeypatch.setattr(run_eval, "predict", flaky)

    class _Args:
        corpus_root = None; tickers = None; limit = None; sample_seed = 0
        tolerance = 1e-4; no_cache = False; api_url = None; cached_only = False
        fail_fast = False; extractor = "pipeline"
    code = run_eval.run_corpus(_Args(), "tag", tmp_path, "corpus_scorecard_tag.json")

    scorecard = json.loads((tmp_path / "corpus_scorecard_tag.json").read_text())
    assert calls["n"] == 2, "the second filing must still have been attempted"
    assert len(scorecard["errors"]) == 2
    assert {e["key"] for e in scorecard["errors"]} == {"ARCI Q1 2022", "ARCI Q2 2022"}
    assert code == 0, "failures are recorded; only a wrong answer fails the run"


def test_the_single_ticker_scorecard_keeps_every_key_the_readme_documents(tmp_path):
    """The corpus additions are additive. These keys are a published contract."""
    import shutil
    run_eval = _script("run_eval")
    shutil.copytree(ROOT / "output" / "predictions", tmp_path / "predictions")

    import os
    os.environ["GEMINI_MODEL"] = "gemini-3.1-flash-lite"
    os.environ["GEMINI_INPUT_MODE"] = "image"
    sys.argv = ["run_eval", "--ticker", "ARCI", "--periods", "Q1", "--out", str(tmp_path)]
    assert run_eval.main() == 0

    scorecard = json.loads(next(tmp_path.glob("scorecard_ARCI_*.json")).read_text())
    assert {"ticker", "tolerance", "run", "periods", "usage", "summary",
            "failure_kinds", "provider", "target"} <= set(scorecard)
    assert set(scorecard["summary"]) >= {"correct", "wrong", "missed", "no_truth",
                                         "both_absent"}
    period = scorecard["periods"]["Q1"]
    assert {"fx_rate", "fields", "elapsed_s", "usage", "abstained"} <= set(period)
    assert {"status", "rel_error", "failure_kind", "pred_idr", "truth_idr", "pred_raw",
            "confidence", "grounded", "abstained"} == set(period["fields"]["aset"])


def test_compare_runs_fails_on_a_token_budget_and_never_on_a_missing_one(tmp_path):
    """A metric absent from one run is not evidence of improvement."""
    compare_runs = _script("compare_runs")
    cells = [_cell().to_dict()]

    def scorecard(tokens):
        return {"provider": "tag", "cells": cells, "summary": {"correct": 1},
                "metrics": {"cost_latency": {
                    "tokens_per_filing": {"input": tokens} if tokens else {},
                    "seconds_per_filing": {}, "calls_per_filing": None}}}

    base = tmp_path / "base.json"; cand = tmp_path / "cand.json"
    base.write_text(json.dumps(scorecard(1000)))
    cand.write_text(json.dumps(scorecard(1500)))
    sys.argv = ["compare_runs", str(base), str(cand), "--max-tokens-increase", "0.10",
                "--no-bootstrap"]
    assert compare_runs.main() == 1, "a 50% token increase is over a 10% budget"

    cand.write_text(json.dumps(scorecard(None)))
    assert compare_runs.main() == 0, "an unrecorded metric must not fail the gate"


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = skipped = 0
    for name, fn in tests:
        if fn.__code__.co_argcount:
            skipped += 1
            print(f"  SKIP  {name}  (needs pytest fixtures)")
            continue
        try:
            fn()
            print(f"  PASS  {name}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL  {name}: {type(exc).__name__}: {exc}")
    print(f"\n{len(tests) - failed - skipped}/{len(tests) - skipped} passed"
          f"{f', {skipped} skipped' if skipped else ''}")
    raise SystemExit(1 if failed else 0)
