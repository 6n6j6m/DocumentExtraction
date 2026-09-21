#!/usr/bin/env python3
"""
The one input to the ratio sheet that is not in the filing: the market price.

Everything else in `data/ground_truth/*.xlsx` is read off a balance sheet. `Harga saham
rupiah` is not — it is what the market paid on a particular day, and it is the numerator
of both PBV and PE. With that row empty, two of the seven ratios cannot be computed at
all, which is why it is worth its own script rather than another column of manual
lookups.

**Which day, and why it is not the period end.** A quarter ends and its report is not
public for weeks. Pricing a Q1 balance sheet at 31 March values it on information nobody
had yet. The convention this script implements takes the price at a date by which the
market has actually seen the numbers:

    Q1 -> 17 June          Q3 -> 17 November
    Q2 -> 17 August        Q4 -> 17 May of the FOLLOWING year

Q4 moves into the next year because the annual report is audited: the December figures
reach the market the following spring, not in December. Pass --q4-same-year if your
convention differs.

**The 17th is often not a trading day.** It falls on a weekend roughly two times in
seven, and Indonesian market holidays take more -- 17 August is Independence Day every
year, so Q2 never lands on an open market. The price used is therefore the close of the
**nearest** trading day, before or after the 17th -- the earlier one when two are equally
near, and never a day that has not happened yet -- and the date actually used is reported
on every row. A silently substituted date is how a price ends up describing a different
week.

**Nothing is invented.** If no trading day is found within the lookback window, the row
is left empty and says so, the same way an abstained field does elsewhere in this repo.

**Close, not adjusted close.** Yahoo's `adjclose` is restated for later splits and
dividends. The sheet divides this price by book value per share and by EPS, both taken
from the filing as it stood at the time, so the price has to be the one that was actually
quoted then. An adjusted price would silently mismatch its own denominator.

Usage:
    # what it would write, from the workbook's own period columns
    python scripts/fetch_share_prices.py --ticker ARCI

    # several issuers, to a CSV
    python scripts/fetch_share_prices.py --ticker ARCI JPFA GTRA -o output/prices.csv

    # explicit periods, no workbook needed
    python scripts/fetch_share_prices.py --ticker CPIN --periods Q1_2024 Q2_2024

    # actually fill the "Harga saham rupiah" row (close Excel first)
    python scripts/fetch_share_prices.py --ticker ARCI --write-xlsx

    # every workbook in data/ground_truth at once; one that fails is skipped, not fatal
    python scripts/fetch_share_prices.py --all --write-xlsx

    # merge into data/share_prices.json instead
    python scripts/fetch_share_prices.py --ticker ARCI --write-json
"""

import argparse
import csv
import json
import re
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
GROUND_TRUTH = ROOT / "data" / "ground_truth"
SHARE_PRICES = ROOT / "data" / "share_prices.json"

# Quarter -> (month, day, years_forward). The forward year on Q4 is the whole reason
# this mapping is data rather than arithmetic on the period end.
QUARTER_DATE = {
    "Q1": (6, 17, 0),
    "Q2": (8, 17, 0),
    "Q3": (11, 17, 0),
    "Q4": (5, 17, 1),
}

# How far back to walk for an open market. A weekend plus Idul Fitri week is the longest
# closure that realistically sits behind the 17th; beyond that, returning nothing is more
# honest than returning a price from a fortnight away.
LOOKBACK_DAYS = 12

CHART_URL = "https://query1.finance.yahoo.com/v8/finance/chart/{symbol}"
# IDX tickers carry the .JK suffix on Yahoo. Kept as a constant because an issuer that
# also lists elsewhere would otherwise silently return the wrong exchange's price.
SUFFIX = ".JK"
HEADERS = {"User-Agent": "Mozilla/5.0"}
TIMEOUT = 20

_PERIOD = re.compile(r"^(Q[1-4])[ _-]?(\d{4})$", re.I)
# The workbooks label the annual column "TAHUNAN 2024", not "Q4 2024" -- every sheet in
# data/ground_truth does, for every year except 2025. Read literally, this script found
# no Q4 column at all for 2020-2024 and silently priced three quarters a year instead of
# four. An annual column IS Q4: same balance-sheet date, same pricing date.
_ANNUAL = re.compile(r"^(?:tahunan|annual|fy)[ _-]?(\d{4})$", re.I)


class PriceUnavailable(Exception):
    """No usable quote. Carries the reason so the row can say what happened."""


def parse_period(text: str):
    """A period header -> ("Q1", 2024). None if it is not one.

    Accepts "Q1_2024" / "Q1 2024" / "q1-2024", and "TAHUNAN 2024" as Q4 2024. Both
    spellings collapse to the same tuple on purpose: the sheet uses one for 2025 and the
    other for every earlier year, and a lookup that told them apart would write the price
    into a column it had not found.
    """
    raw = str(text).strip()
    match = _PERIOD.match(raw)
    if match:
        return (match.group(1).upper(), int(match.group(2)))
    annual = _ANNUAL.match(raw)
    return ("Q4", int(annual.group(1))) if annual else None


def price_date(quarter: str, year: int, q4_same_year: bool = False) -> date:
    """The date whose close prices this period."""
    month, day, forward = QUARTER_DATE[quarter]
    if quarter == "Q4" and q4_same_year:
        forward = 0
    return date(year + forward, month, day)


def fetch_daily_closes(ticker: str, start: date, end: date) -> dict:
    """{date: close} for one symbol over a window. Raises PriceUnavailable.

    One request per period is wasteful only in theory: the periods of one issuer are
    months apart, so a single window covering all of them would download years of daily
    bars to use eight of them.
    """
    symbol = ticker if "." in ticker else ticker.upper() + SUFFIX
    params = {
        # Yahoo treats the range as half-open and in exchange-local terms; a day of
        # padding on each side costs nothing and avoids losing the boundary bar.
        "period1": int(datetime(start.year, start.month, start.day,
                                tzinfo=timezone.utc).timestamp()) - 86400,
        "period2": int(datetime(end.year, end.month, end.day,
                                tzinfo=timezone.utc).timestamp()) + 86400,
        "interval": "1d",
    }
    try:
        response = requests.get(CHART_URL.format(symbol=symbol), params=params,
                                headers=HEADERS, timeout=TIMEOUT)
    except requests.RequestException as exc:
        raise PriceUnavailable(f"network error: {exc}") from exc

    if response.status_code == 404:
        raise PriceUnavailable(f"unknown symbol {symbol}")
    if response.status_code != 200:
        # Yahoo answers 400 with "Data doesn't exist for startDate..." for a window that
        # predates the listing, which is the single most likely reason to land here and
        # means something quite different from a broken request. ARCI listed in June
        # 2021, so its 2020 columns can never be priced, and the row should say that
        # rather than quote a status code at somebody.
        detail = ""
        try:
            detail = (response.json().get("chart", {}).get("error") or {}).get(
                "description", "")
        except ValueError:
            pass
        if response.status_code == 400 and "Data doesn't exist" in detail:
            raise PriceUnavailable(
                f"{symbol} has no trading data around {end} "
                f"(listed later, delisted, or suspended)")
        raise PriceUnavailable(f"HTTP {response.status_code} from Yahoo"
                               + (f": {detail}" if detail else ""))

    try:
        result = response.json()["chart"]["result"][0]
    except (ValueError, KeyError, TypeError, IndexError):
        raise PriceUnavailable(f"no price series returned for {symbol}")

    meta = result.get("meta") or {}
    currency = (meta.get("currency") or "").upper()
    if currency and currency != "IDR":
        # The sheet's column is named "Harga saham rupiah" and divides rupiah book value.
        # A price in another currency would be off by the exchange rate and look plausible.
        raise PriceUnavailable(f"{symbol} quotes in {currency}, not IDR")

    stamps = result.get("timestamp") or []
    closes = (result.get("indicators", {}).get("quote") or [{}])[0].get("close") or []
    # Exchange-local, not UTC: a Jakarta session opens at 02:00 UTC, so reading the
    # timestamp in UTC happens to work -- until the offset or the open time changes.
    offset = meta.get("gmtoffset", 0) or 0

    series = {}
    for stamp, close in zip(stamps, closes):
        if close is None:
            continue            # a halted or suspended session; not a zero price
        series[datetime.fromtimestamp(stamp + offset, tz=timezone.utc).date()] = float(close)
    if not series:
        raise PriceUnavailable(f"no trading days for {symbol} in the window")
    return series


def close_nearest(series: dict, target: date, window: int = LOOKBACK_DAYS,
                  today: date = None):
    """(date, close) for the open session nearest `target`, either side of it.

    The 17th is a convention, not a hard date, so a session two days after it prices the
    period as well as one two days before. At equal distance the earlier one wins -- it is
    the price the market set without the extra days of news. A session after `today` does
    not exist yet and is never used.
    """
    today = today or date.today()
    for distance in range(window + 1):
        for day in (target - timedelta(days=distance), target + timedelta(days=distance)):
            if day <= today and day in series:
                return day, series[day]
    raise PriceUnavailable(f"market closed for {window}+ days either side of {target}")


def periods_from_workbook(ticker: str) -> list:
    """The period columns the sheet actually has, newest first, as it orders them.

    Taken from the workbook rather than from the filenames so the value lands in the
    column it was looked up for. A list built independently would drift the day someone
    inserts a quarter.
    """
    import openpyxl

    path = GROUND_TRUTH / f"{ticker.upper()}.xlsx"
    if not path.exists():
        raise FileNotFoundError(f"no ground-truth workbook at {path}")
    worksheet = openpyxl.load_workbook(path, read_only=True).active
    for row in worksheet.iter_rows(min_row=1, max_row=12, values_only=True):
        found = [parse_period(cell) for cell in row if isinstance(cell, str)]
        found = [f for f in found if f]
        if len(found) >= 2:
            return found
    raise ValueError(f"{path.name}: no row of period headers (\"Q1 2024\" ...) found")


def periods_from_filings(directory: Path, ticker: str) -> list:
    """Periods discovered from <Quarter>_<Year>_<TICKER>.pdf, the repo's own convention."""
    found = []
    for pdf in sorted(directory.glob(f"*_{ticker.upper()}.pdf")):
        parsed = parse_period("_".join(pdf.stem.split("_")[:2]))
        if parsed:
            found.append(parsed)
    return sorted(set(found), key=lambda p: (p[1], p[0]), reverse=True)


def collect(ticker: str, periods: list, q4_same_year: bool = False) -> list:
    """One row per period: the date used, the close, or the reason there is neither."""
    rows = []
    for quarter, year in periods:
        target = price_date(quarter, year, q4_same_year)
        row = {"ticker": ticker.upper(), "period": f"{quarter} {year}",
               "key": f"{ticker.upper()}_{year}{quarter}",
               "target_date": target.isoformat(), "price_date": "", "price": "",
               "days_back": "", "note": ""}
        if target > date.today():
            # A period whose pricing date has not arrived yet. Not an error -- the sheet
            # simply cannot have that column filled in yet, and saying so beats a blank.
            row["note"] = f"pricing date {target} is in the future"
            rows.append(row)
            continue
        try:
            series = fetch_daily_closes(ticker, target - timedelta(days=LOOKBACK_DAYS),
                                        min(target + timedelta(days=LOOKBACK_DAYS),
                                            date.today()))
            used, close = close_nearest(series, target)
            # Positive: the session was before the 17th; negative: after it.
            row.update(price_date=used.isoformat(), price=close,
                       days_back=(target - used).days)
            if used != target:
                side = "previous" if used < target else "next"
                row["note"] = f"market closed on {target}; used the nearest ({side}) session"
        except PriceUnavailable as exc:
            row["note"] = str(exc)
        rows.append(row)
        time.sleep(0.3)         # a courtesy pause; this is somebody's free endpoint
    return rows


def write_json(rows: list, path: Path = SHARE_PRICES) -> int:
    """Merge into data/share_prices.json, keyed <TICKER>_<YYYY>Q<n> as it already is."""
    data = json.loads(path.read_text()) if path.exists() else {}
    written = 0
    for row in rows:
        if row["price"] != "":
            data[row["key"]] = str(int(row["price"]) if float(row["price"]).is_integer()
                                   else row["price"])
            written += 1
    # Keys sorted, but the file's own explanation kept at the top: `_comment` sorts
    # after every uppercase ticker, so plain sort_keys buries the one line that tells
    # the next reader where these numbers came from.
    comment = {k: data.pop(k) for k in list(data) if k.startswith("_")}
    ordered = {**comment, **{k: data[k] for k in sorted(data)}}
    path.write_text(json.dumps(ordered, indent=2) + "\n")
    return written


def write_xlsx(ticker: str, rows: list, overwrite: bool = False) -> int:
    """Fill the "Harga saham rupiah" row, and nothing else.

    Only empty cells are filled unless `overwrite` is set. A price already in the sheet
    was put there by someone, and it can disagree with Yahoo: TLDN's sheet has 472 for
    Q1 2026 where Yahoo closes 530. Which one is right is a decision about ground truth,
    so the disagreement is printed and the cell is left alone.

    Located by label and by header text rather than by row and column number, so a sheet
    with an inserted row still lands the value in the right cell. Refuses rather than
    guesses if either cannot be found: writing a price into the wrong row would corrupt
    ground truth in a way that looks like a bad extraction later.
    """
    import openpyxl

    path = GROUND_TRUTH / f"{ticker.upper()}.xlsx"
    lock = path.with_name("~$" + path.name)
    if lock.exists():
        raise RuntimeError(f"{path.name} looks open in Excel ({lock.name} exists). "
                           f"Close it first -- Excel will overwrite whatever is written "
                           f"here when it saves.")

    workbook = openpyxl.load_workbook(path)
    worksheet = workbook.active

    price_row = None
    columns = {}
    for row in worksheet.iter_rows(min_row=1, max_row=40):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            if cell.value.strip().lower().startswith("harga saham"):
                price_row = cell.row
            parsed = parse_period(cell.value)
            if parsed:
                columns[parsed] = cell.column

    if price_row is None:
        raise RuntimeError(f"{path.name}: no row labelled 'Harga saham ...' found")
    if not columns:
        raise RuntimeError(f"{path.name}: no period column headers found")

    written = 0
    for row in rows:
        parsed = parse_period(row["period"])
        if row["price"] == "" or parsed not in columns:
            continue
        cell = worksheet.cell(row=price_row, column=columns[parsed])
        existing = cell.value
        if existing not in (None, ""):
            try:
                agrees = abs(float(existing) - float(row["price"])) < 0.5
            except (TypeError, ValueError):
                agrees = False
            if agrees:
                continue
            if not overwrite:
                print(f"    ≠ {row['period']}: sheet has {existing}, Yahoo closed "
                      f"{row['price']:,.0f} on {row['price_date']} - left as is "
                      f"(--overwrite to replace)")
                continue
            print(f"    ↺ {row['period']}: {existing} replaced by {row['price']:,.0f}")
        cell.value = row["price"]
        written += 1
    if written:
        workbook.save(path)
    return written


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Fetch the closing share price that prices each quarter.")
    who = ap.add_mutually_exclusive_group(required=True)
    who.add_argument("--ticker", nargs="+", help="one or more IDX tickers")
    who.add_argument("--all", action="store_true",
                     help="every workbook in data/ground_truth (one ticker per <TICKER>.xlsx)")
    ap.add_argument("--periods", nargs="+",
                    help="explicit periods (Q1_2024 ...); default: read the workbook")
    ap.add_argument("--from-filings", metavar="DIR",
                    help="discover periods from <Q>_<YYYY>_<TICKER>.pdf in DIR instead")
    ap.add_argument("--q4-same-year", action="store_true",
                    help="price Q4 on 17 May of the SAME year (default: the next one)")
    ap.add_argument("-o", "--out", help="also write the table to this CSV")
    ap.add_argument("--write-json", action="store_true",
                    help=f"merge prices into {SHARE_PRICES.relative_to(ROOT)}")
    ap.add_argument("--write-xlsx", action="store_true",
                    help="fill the 'Harga saham rupiah' row in the ground-truth workbook")
    ap.add_argument("--overwrite", action="store_true",
                    help="with --write-xlsx, also replace prices already in the sheet "
                         "that differ from Yahoo (default: fill empty cells only)")
    args = ap.parse_args()

    tickers = args.ticker or sorted(
        p.stem for p in GROUND_TRUTH.glob("*.xlsx") if not p.name.startswith(("~$", "_")))
    if args.periods and any(not parse_period(x) for x in args.periods):
        bad = [x for x in args.periods if not parse_period(x)]
        print(f"not periods: {', '.join(bad)} (expected Q1_2024)")
        return 2

    # One workbook that cannot be read or written (open in Excel, no price row) is
    # reported and skipped. Stopping there would leave every ticker after it unpriced,
    # which is the opposite of what a batch over all workbooks is for.
    all_rows, failed = [], []
    for ticker in tickers:
        if args.periods:
            periods = [parse_period(x) for x in args.periods]
        elif args.from_filings:
            periods = periods_from_filings(Path(args.from_filings), ticker)
        else:
            try:
                periods = periods_from_workbook(ticker)
            except (FileNotFoundError, ValueError) as exc:
                print(f"✗ {ticker}: {exc}")
                print("  pass --periods Q1_2024 ... or --from-filings DIR")
                failed.append(ticker)
                continue

        if not periods:
            print(f"✗ {ticker}: no periods found")
            failed.append(ticker)
            continue

        print(f"\n{ticker.upper()} — {len(periods)} period(s)")
        rows = collect(ticker, periods, args.q4_same_year)
        print(f"  {'period':10} {'priced on':12} {'used':12} {'close':>10}  note")
        for row in rows:
            print(f"  {row['period']:10} {row['target_date']:12} "
                  f"{row['price_date'] or '-':12} "
                  f"{(f'{row['price']:,.0f}' if row['price'] != '' else '-'):>10}  "
                  f"{row['note']}")
        all_rows += rows

        if args.write_xlsx:
            try:
                print(f"  → {write_xlsx(ticker, rows, args.overwrite)} cell(s) written to "
                      f"{ticker.upper()}.xlsx")
            except RuntimeError as exc:
                print(f"  ✗ {exc}")
                failed.append(ticker)

    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        with out.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(all_rows[0]))
            writer.writeheader()
            writer.writerows(all_rows)
        print(f"\n{len(all_rows)} row(s) → {out}")

    if args.write_json:
        print(f"{write_json(all_rows)} price(s) merged into "
              f"{SHARE_PRICES.relative_to(ROOT)}")

    missing = [r for r in all_rows if r["price"] == ""]
    if missing:
        print(f"\n{len(missing)} period(s) have no price:")
        for row in missing:
            print(f"  {row['ticker']} {row['period']}: {row['note']}")
    if failed:
        print(f"\n{len(failed)} workbook(s) skipped: {', '.join(failed)}")
    # 0 every period priced, 1 some blank or a workbook skipped. A blank is a fact, not a
    # crash, but a caller scripting this needs to know without parsing the output.
    return 0 if not missing and not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
