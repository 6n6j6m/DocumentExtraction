#!/usr/bin/env python3
"""
Evaluate the extraction pipeline against the ground truth in data/ground_truth/<TICKER>.xlsx.

The spreadsheet is the system of record: one column per period, one row per field
(the row numbers live in src/schema.py EXCEL_ROWS). Its figures are stated in full
Rupiah, so predictions -- which come out of the filing in whatever currency and
scale it was printed in -- are normalised to Rupiah before being compared.

Two comparisons are reported per field, because they fail for different reasons:

  as-printed (USD)  isolates EXTRACTION error -- did the model read the right line?
  converted (IDR)   adds CONVERSION error    -- was the scale and rate right?

A field can be right in USD and wrong in IDR (bad rate/scale), which is a different
bug from reading the wrong row, so they are never collapsed into one number.

Predictions are cached under output/predictions/, so re-running to inspect results
costs no API calls. Use --no-cache to force a fresh extraction.

Usage:
    python scripts/run_eval.py                          # all periods
    python scripts/run_eval.py --periods Q1 Q2          # a subset
    python scripts/run_eval.py --no-cache --tolerance 0.001
"""

import argparse
import hashlib
import json
import os
import random
import subprocess
import time
from datetime import datetime, timezone
import sys
from pathlib import Path
from typing import Optional

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

from schema import FinancialStatementExtraction, EXCEL_ROWS
from normalize import to_idr, extract_fx_rate
from extract import extract_from_pdf
from llm import get_provider_with_fallback, LLMError, LLMUnavailable
from usage import Usage
from periods import QUARTER_END, parse_period, period_label
# Scoring lives in src/evalkit.py so the metrics can be exercised without a model, a
# key or a PDF; this script is the driver. compare() and classify_failure() are
# imported rather than redefined -- two copies of the status vocabulary is how a
# scorecard and its comparison tool drift apart.
from evalkit import (Cell, DerivedInputs, PeriodRef, cache_path_for, classify_failure,
                     compare, conversion_sha, derive, derive_from_cache, discover_corpus,
                     legacy_cache_path, load_truth_sheet, score_period)
from evalmetrics import bootstrap_delta, summarise

HEADER_ROW = 3

# Environment that changes what the run measures. Recorded with the scorecard so a
# difference between two runs can be attributed instead of assumed.
TRACKED_CONFIG = [
    "LLM_PROVIDER", "GEMINI_MODEL", "OLLAMA_MODEL",
    "GEMINI_INPUT_MODE", "OLLAMA_INPUT_MODE",
    "SPLIT_BY_STATEMENT", "CONDENSE_TEXT", "IMAGE_DPI",
    "CONFIDENCE_ABSTAIN_THRESHOLD", "MAX_PDF_PAGES",
    # Which implementation answered, and what it was allowed to spend. Without these in
    # the run's config, compare_runs would diff a pipeline scorecard against an agent
    # one and report the difference as a regression in the extractor.
    "EXTRACTOR", "AGENT_MAX_STEPS", "AGENT_MAX_IMAGES", "AGENT_MAX_TEXT_PAGES",
]


def run_metadata(tolerance: float) -> dict:
    """Everything needed to say whether two scorecards are comparable.

    A score that moved between runs is only informative if you know what else moved.
    The prompt hash matters most: prompts are edited far more often than code, and a
    prompt change is invisible in a git diff of the harness.
    """
    from prompts import SYSTEM_PROMPT, EXTRACTION_PROMPT

    def git(*args):
        try:
            return subprocess.check_output(["git", *args], cwd=ROOT,
                                           stderr=subprocess.DEVNULL).decode().strip()
        except Exception:
            # No .git -- running inside the image, which excludes it. The build stamps
            # GIT_COMMIT so the scorecard can still say which code produced it; without
            # that, a containerised run is unattributable and cannot be compared with
            # anything.
            return os.getenv("GIT_COMMIT") if args[:1] == ("rev-parse",) else None

    # Tracked modifications only. An untracked scratch file in data/ does not make a
    # run irreproducible, but a modified source file does -- lumping the two together
    # leaves the flag permanently true and therefore meaningless.
    dirty = git("status", "--porcelain", "--untracked-files=no")
    return {
        "timestamp": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "git_commit": git("rev-parse", "--short", "HEAD"),
        "git_dirty": bool(dirty) if dirty is not None else None,
        "prompt_sha256": hashlib.sha256(
            (SYSTEM_PROMPT + EXTRACTION_PROMPT).encode()).hexdigest()[:12],
        "tolerance": tolerance,
        "config": {k: os.getenv(k) for k in TRACKED_CONFIG if os.getenv(k) is not None},
    }


def discover_periods(ticker: str) -> dict:
    """Find the filings present for a ticker and map them to spreadsheet columns.

    Derived from what is on disk rather than hard-coded, so a second issuer needs no
    code change. Filenames follow <Quarter>_<Year>_<TICKER>.pdf.

    The Q4 filing is the annual report, which the spreadsheet heads "TAHUNAN <year>"
    rather than "Q4 <year>" -- the figures are cumulative for the year, so the label
    reflects that.
    """
    periods = {}
    for pdf in sorted((ROOT / "data" / "raw").glob(f"*_{ticker}.pdf")):
        parts = pdf.stem.split("_")
        if len(parts) < 3 or parts[0] not in QUARTER_END:
            continue
        quarter, year = parts[0], parts[1]
        label = f"TAHUNAN {year}" if quarter == "Q4" else f"{quarter} {year}"
        periods[quarter] = (label, pdf.stem, f"{year}-{QUARTER_END[quarter]}")
    return periods


TICKER_CELL = "D1"      # the sheet names its own issuer here


def load_ground_truth(xlsx_path: Path, period_label: str, ticker: str = None) -> dict:
    """Read one period's column. Returns {field: idr_value_or_None}.

    The sheet is checked against the ticker being evaluated first. These workbooks
    are made by copying an existing one and overwriting the columns, so a copy that
    was never re-labelled scores one issuer's extraction against another issuer's
    figures -- and every field comes back `wrong` for a reason that is nowhere in
    the output. The sheet names its own issuer; comparing costs one cell read.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Sheet1"]

    sheet_ticker = ws[TICKER_CELL].value
    if ticker and isinstance(sheet_ticker, str) and sheet_ticker.strip().upper() != ticker.upper():
        raise ValueError(
            f"{xlsx_path.name} is labelled for {sheet_ticker.strip()!r} in cell "
            f"{TICKER_CELL}, but the run is for {ticker!r}. This is what an "
            f"un-relabelled copy of another issuer's sheet looks like; fix the sheet "
            f"rather than scoring against the wrong figures."
        )

    column = None
    for cell in ws[HEADER_ROW]:
        if isinstance(cell.value, str) and cell.value.strip() == period_label:
            column = cell.column_letter
            break
    if column is None:
        raise KeyError(f"No column headed {period_label!r} in row {HEADER_ROW} of {xlsx_path.name}")

    truth = {}
    for field, (row, _label) in EXCEL_ROWS.items():
        value = ws[f"{column}{row}"].value
        truth[field] = float(value) if isinstance(value, (int, float)) else None
    return truth


def extract_via_api(pdf_path: Path, api_url: str, timeout: int = 600):
    """Extract by POSTing the filing to a running service.

    The only difference from the in-process path is *where the extraction happens*.
    Everything after this -- normalising to Rupiah, comparing, classifying failures --
    is the same code operating on the same object, which is what makes the two paths
    comparable by construction rather than by inspection. A second scoring
    implementation for the service would be a second thing to keep correct, and the
    scorecard could not then be used to say whether the service and the library agree.

    Raises LLMError so the caller's existing per-period recovery applies unchanged.
    """
    import httpx

    url = api_url.rstrip("/") + "/extract"
    try:
        with open(pdf_path, "rb") as handle:
            response = httpx.post(
                url,
                files={"file": (pdf_path.name, handle, "application/pdf")},
                timeout=timeout,
            )
    except httpx.HTTPError as exc:
        raise LLMError(f"could not reach {url}: {exc}")

    if response.status_code == 503:
        # The service is up but its provider is not. Same meaning as the in-process
        # LLMUnavailable, and the caller already knows what to do with it.
        raise LLMError(f"API reports provider unavailable: "
                       f"{response.json().get('detail', response.text)[:200]}")
    if response.status_code != 200:
        raise LLMError(f"API {response.status_code}: {response.text[:200]}")

    body = response.json()
    result = FinancialStatementExtraction.from_dict(body["extraction"])
    result.field_confidence = body.get("confidence", {})
    result.usage_dict = body.get("usage")
    return result


def predict(pdf_path: Path, cache_path: Path, use_cache: bool,
            api_url: Optional[str] = None, legacy_path: Optional[Path] = None,
            extractor: str = "pipeline"):
    """Extract from a filing, caching the raw (as-printed) result.

    `api_url` chooses WHERE extraction runs, not how it is scored.

    Returns (result, cached, payload). The payload is the cache dict itself, so the
    caller can read and backfill the conversion cached beside the prediction.
    """
    source = cache_path if (use_cache and cache_path.exists()) else (
        legacy_path if use_cache else None)
    if source is not None and source.exists():
        data = json.loads(source.read_text())
        # Cached payloads are read back through the same type checks as a live model
        # response: a hand-edited or stale cache is exactly the kind of input that used
        # to raise deep inside a validation rule instead of being rejected here.
        result = FinancialStatementExtraction.from_dict(
            {k: v for k, v in data.items() if not k.startswith("_")})
        # Confidence lives beside the fields rather than inside the dataclass, so it
        # has to be restored explicitly or a cached run loses it silently.
        result.field_confidence = data.get("_confidence", {})
        result.elapsed_s = data.get("_elapsed_s")
        result.usage_dict = data.get("_usage")
        result.model = data.get("_model")
        # Without this, a re-scored agent run silently loses every agent metric -- the
        # same failure _confidence already carries a comment about.
        result.agent_trajectory = data.get("_agent")
        return result, True, data

    started = time.time()
    if api_url:
        result = extract_via_api(pdf_path, api_url)
    else:
        usage = Usage()
        if extractor == "agentic":
            # Imported here, not at module scope: a syntax error in new agent code must
            # not cost the committed baseline its ability to re-score.
            from agent import extract_agentic
            result = extract_agentic(str(pdf_path), usage=usage)
        else:
            result = extract_from_pdf(str(pdf_path), usage=usage)
        if result is not None:
            result.usage_dict = usage.to_dict()
    if result is None:
        return None, False, {}
    result.elapsed_s = round(time.time() - started, 1)

    payload = result.to_dict()
    payload["_confidence"] = getattr(result, "field_confidence", {})
    payload["_elapsed_s"] = result.elapsed_s
    payload["_usage"] = getattr(result, "usage_dict", None)
    # Which model actually answered. With a fallback chain a throttled filing is handed
    # to another model, and a run that does not record that reports one number for two
    # systems -- so `model` is a slice in the metrics, not a footnote.
    payload["_model"] = getattr(result, "model", None)
    payload["_agent"] = getattr(result, "agent_trajectory", None)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(json.dumps(payload, indent=2))
    return result, False, payload


def derived_for(pred, pdf: Path, cache_path: Path, payload: dict) -> DerivedInputs:
    """The conversion to full Rupiah, from cache when it is still valid.

    Scoring used to re-open the filing every time -- page selection, text, FX rate --
    even when the prediction came from cache. On EMAS, whose text layer is glyph ids,
    that is minutes of OCR to re-derive an answer nobody re-asked the model for. The
    result is cached beside the prediction and backfilled here, which is what lets CI
    score the corpus with no PDFs, no poppler and no key.
    """
    cached = derive_from_cache(payload)
    if cached is not None:
        return cached
    derived = derive(pred, pdf)
    if payload:
        payload["_derived"] = derived.to_dict()
        try:
            cache_path.parent.mkdir(parents=True, exist_ok=True)
            cache_path.write_text(json.dumps(payload, indent=2, default=str))
        except OSError:
            pass          # a read-only cache is not a reason to lose the run
    return derived


def provider_tag_for(provider, api_url=None, extractor: str = "pipeline") -> str:
    """The name a run is filed under: model, input mode, target, implementation.

    Every element is here because leaving it out made two different experiments share a
    scorecard and a prediction cache -- and a comparison that is supposed to settle a
    question quietly became a run compared with itself.

      input mode   the model is shown different things; that is a different experiment
      _api         the service and the library are two systems, same model or not
      _agentic     a tool-calling loop is a different system again, on the same model
    """
    tag = (f"{provider.name}_{getattr(provider, 'input_mode', 'text')}"
           .replace(":", "_").replace("/", "_"))
    if api_url:
        tag += "_api"
    if extractor and extractor != "pipeline":
        tag += f"_{extractor}"
    return tag


GROUND_TRUTH = ROOT / "data" / "ground_truth"
FLUSH_EVERY = 10        # a killed three-hour run still leaves a readable scorecard


def corpus_roots(explicit) -> list:
    """Where filings are looked for.

    `data/raw` is the default and the only one a reviewer needs: it is in the repo, so
    CI and a fresh checkout score the same filings. A personal archive
    (~/FinancialReport/<TICKER>/...) is opted into with --corpus-root or by setting
    EVAL_CORPUS_ROOT, never assumed, because a scorecard that silently depends on files
    nobody else has is not a shared measurement.
    """
    if explicit:
        return [Path(r).expanduser() for r in explicit]
    roots = [ROOT / "data" / "raw"]
    extra = os.getenv("EVAL_CORPUS_ROOT", "")
    roots += [Path(r).expanduser() for r in extra.split(os.pathsep) if r.strip()]
    return roots


def corpus_pairs(roots, tickers, limit=None, seed=0):
    """Every (ticker, period) with BOTH a filing and a labelled column.

    Returns (pairs, sheets, skipped). A workbook that cannot be trusted for its ticker
    is named in `skipped` and its filings are not scored -- a corpus run reports the
    problem and carries on, where a single-ticker run still raises.
    """
    found = discover_corpus(roots, tickers)
    pairs, sheets, skipped = [], {}, []
    for ticker in sorted(found):
        sheet = load_truth_sheet(GROUND_TRUTH / f"{ticker}.xlsx", ticker)
        if sheet.problem:
            skipped.append({"ticker": ticker, "reason": sheet.problem,
                            "filings": len(found[ticker])})
            continue
        sheets[ticker] = sheet
        for (quarter, year), pdf in found[ticker].items():
            column = sheet.columns.get((quarter, year))
            if not column or all(v is None for v in column.values()):
                continue            # a filing nobody has labelled yet is not a miss
            pairs.append(PeriodRef(ticker=ticker, quarter=quarter, year=year, pdf=pdf))
    pairs.sort(key=lambda ref: ref.order)
    if limit and limit < len(pairs):
        # A sampled pass is for checking the machinery cheaply, so the sample is
        # reproducible: the same seed scores the same filings.
        pairs = sorted(random.Random(seed).sample(pairs, limit), key=lambda r: r.order)
    return pairs, sheets, skipped


def run_corpus(args, provider_tag, out_dir, scorecard_name) -> int:
    """Score every labelled filing, not one issuer's four quarters.

    The prediction cache is the resume mechanism: a re-run skips what is already
    extracted, so a pass interrupted after two hours costs two hours less the second
    time. Every failure is recorded and stepped over, because in a 223-filing run the
    alternative is one broken PDF ending the measurement.
    """
    roots = corpus_roots(args.corpus_root)
    pairs, sheets, skipped = corpus_pairs(roots, args.tickers, args.limit, args.sample_seed)
    labelled_cells = sum(sheets[r.ticker].columns[(r.quarter, r.year)] is not None
                         and sum(1 for v in sheets[r.ticker].columns[(r.quarter, r.year)].values()
                                 if v is not None)
                         for r in pairs)

    print(f"corpus: {len(pairs)} filing(s), {labelled_cells} labelled cell(s), "
          f"from {', '.join(str(r) for r in roots)}")
    for entry in skipped:
        print(f"  ! {entry['ticker']} skipped: {entry['reason']}")
    if not pairs:
        print("nothing to score -- no filing has a labelled column")
        return 2

    report = {"ticker": "CORPUS", "tolerance": args.tolerance,
              "run": run_metadata(args.tolerance), "periods": {}}
    tally = {"correct": 0, "wrong": 0, "missed": 0, "no_truth": 0, "both_absent": 0}
    failure_kinds, errors, all_cells, filings = {}, [], [], []
    trajectories = []
    uncached = []

    def flush():
        report["summary"] = tally
        report["failure_kinds"] = failure_kinds
        report["provider"] = provider_tag
        report["target"] = args.api_url or "in-process"
        report["schema_version"] = 2
        report["cells"] = [c.to_dict() for c in all_cells]
        report["metrics"] = summarise(all_cells, filings, trajectories)
        report["conversion_sha"] = conversion_sha()
        report["usage"] = report["metrics"]["cost_latency"]
        report["corpus"] = {
            "roots": [str(r) for r in roots],
            "tickers": sorted({r.ticker for r in pairs}),
            "pairs": len(pairs), "cells_labelled": labelled_cells,
            "scored_pairs": len({c.key for c in all_cells}),
            "skipped": skipped, "uncached": uncached,
            "models_used": dict(sorted(
                {m: sum(1 for c in all_cells if c.model == m and c.field == "aset")
                 for m in {c.model for c in all_cells}}.items(), key=lambda kv: -kv[1])),
        }
        report["errors"] = errors
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / scorecard_name).write_text(json.dumps(report, indent=2, default=str))

    for index, ref in enumerate(pairs, 1):
        cache_path = cache_path_for(ref, out_dir, provider_tag)
        legacy = legacy_cache_path(ref, out_dir, provider_tag)
        if args.cached_only and not cache_path.exists() and legacy is None:
            uncached.append(ref.key)
            continue

        truth = sheets[ref.ticker].columns[(ref.quarter, ref.year)]
        try:
            pred, cached, payload = predict(ref.pdf, cache_path, not args.no_cache,
                                            api_url=args.api_url, legacy_path=legacy,
                                            extractor=args.extractor)
            if pred is None:
                raise LLMError("extraction returned no fields")
            derived = derived_for(pred, ref.pdf, cache_path, payload)
            cells = score_period(ref, pred, truth, derived, args.tolerance)
        except Exception as exc:              # one filing, not the corpus
            print(f"[{index}/{len(pairs)}] {ref.key:20} FAILED  {type(exc).__name__}: {exc}")
            errors.append({"key": ref.key, "error": f"{type(exc).__name__}: {exc}"})
            report["periods"][ref.key] = {"error": str(exc)}
            if args.fail_fast:
                flush()
                return 1
            continue

        rows = {}
        for cell in cells:
            tally[cell.status] += 1
            if cell.failure_kind:
                failure_kinds[cell.failure_kind] = failure_kinds.get(cell.failure_kind, 0) + 1
            rows[cell.field] = {"status": cell.status, "rel_error": cell.rel_error,
                                "failure_kind": cell.failure_kind,
                                "pred_idr": cell.pred_idr, "truth_idr": cell.truth_idr,
                                "pred_raw": cell.pred_raw, "confidence": cell.confidence,
                                "grounded": cell.grounded, "abstained": cell.abstained}
        all_cells += cells
        filings.append({"elapsed_s": getattr(pred, "elapsed_s", None),
                        "usage": getattr(pred, "usage_dict", None)})
        trajectory = getattr(pred, "agent_trajectory", None)
        if trajectory is not None:
            # The filing key travels with the trajectory: page_selection_agreement names
            # the filings whose pages differ, and a list of page numbers with no filing
            # attached cannot be looked at afterwards.
            trajectory.setdefault("key", ref.key)
        trajectories.append(trajectory)
        report["periods"][ref.key] = {
            "fx_rate": derived.fx_rate, "fields": rows,
            "elapsed_s": getattr(pred, "elapsed_s", None),
            "usage": getattr(pred, "usage_dict", None),
            "abstained": [f for f, r in rows.items() if r.get("abstained")]}

        scored = [c for c in cells if c.labelled]
        right = sum(1 for c in scored if c.status == "correct")
        print(f"[{index}/{len(pairs)}] {ref.key:20} {right}/{len(scored)}"
              f"{'  (cached)' if cached else ''}"
              f"{'  ' + derived.model if derived.model else ''}")
        if index % FLUSH_EVERY == 0:
            flush()

    flush()
    metrics = report["metrics"]
    accuracy = metrics["accuracy"]
    coverage = metrics["coverage"]
    print(f"\n{'='*78}\nCORPUS SUMMARY")
    print(f"  filings scored     {report['corpus']['scored_pairs']}/{len(pairs)}")
    if uncached:
        print(f"  not cached         {len(uncached)}  (--cached-only; run without it to extract)")
    if errors:
        print(f"  failed             {len(errors)}")
    print(f"  cells labelled     {coverage['labelled']}")
    print(f"  correct            {accuracy['correct']}")
    print(f"  wrong              {accuracy['wrong']}")
    print(f"  missed             {accuracy['missed']}  "
          f"(of which withheld deliberately: {coverage['abstained_on_labelled']})")
    if accuracy["accuracy"] is not None:
        print(f"  accuracy           {accuracy['accuracy']:.1%}"
              f"   answered {coverage['answer_rate']:.1%}"
              f"   accuracy when answered {accuracy['accuracy_on_answered']:.1%}")
    withdrawal = metrics["abstention"]
    if withdrawal["withdrawn"]:
        precision = withdrawal["precision"]
        print(f"  abstention         {withdrawal['withdrawn']} withdrawn, "
              f"precision {f'{precision:.0%}' if precision is not None else '-'}"
              f" ({withdrawal['caught']} caught, {withdrawal['thrown_away']} thrown away,"
              f" {withdrawal['unknown_withheld']} unknown)")
    if failure_kinds:
        print("\n  failures by kind:")
        for kind, count in sorted(failure_kinds.items(), key=lambda kv: -kv[1]):
            print(f"    {kind:28}{count:>4}")
    print(f"\n  written to {out_dir / scorecard_name}")
    print(f"  report:  python scripts/report_eval.py {out_dir / scorecard_name}")
    return 1 if tally["wrong"] else 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ticker", default="ARCI")
    ap.add_argument("--periods", nargs="*", default=None)
    ap.add_argument("--tolerance", type=float, default=1e-4,
                    help="relative tolerance, default 0.0001 (0.01%%)")
    ap.add_argument("--no-cache", action="store_true")
    ap.add_argument("--out", default="output")
    ap.add_argument("--provider", help="override LLM_PROVIDER for this run (gemini|ollama)")
    ap.add_argument("--api-url", default=os.getenv("EVAL_TARGET") or None,
                    help="score a running service instead of the library "
                         "(e.g. http://localhost:8000). Defaults to $EVAL_TARGET; "
                         "empty means in-process, so the tests need no container.")
    ap.add_argument("--extractor", choices=["pipeline", "agentic"],
                    default=os.getenv("EXTRACTOR", "pipeline"),
                    help="which implementation to measure: the deterministic pipeline, "
                         "or the tool-calling agent over the same filings, scored by the "
                         "same code")
    ap.add_argument("--corpus", action="store_true",
                    help="score every (ticker, period) that has a filing AND a label, "
                         "instead of one ticker")
    ap.add_argument("--tickers", nargs="+", help="restrict --corpus to these issuers")
    ap.add_argument("--corpus-root", nargs="+",
                    help="where filings live (default: data/raw, plus $EVAL_CORPUS_ROOT)")
    ap.add_argument("--cached-only", action="store_true",
                    help="never call a provider: filings with no cached prediction are "
                         "reported as uncached and skipped. This is what makes a CI run "
                         "keyless by construction rather than by hope.")
    ap.add_argument("--limit", type=int, help="score a random sample of this many filings")
    ap.add_argument("--sample-seed", type=int, default=0,
                    help="seed for --limit, so a sampled pass is reproducible")
    ap.add_argument("--fail-fast", action="store_true",
                    help="stop at the first failing filing instead of recording it")
    ap.add_argument("--allow-fallback", action=argparse.BooleanOptionalAction, default=None,
                    help="keep the provider's fallback chain. Default: on for --corpus "
                         "(it measures the system as deployed, and every cell records "
                         "which model answered), off otherwise (one scorecard, one model)")
    args = ap.parse_args()
    # Back into the env before run_metadata() reads it, so the scorecard records which
    # implementation produced it.
    os.environ["EXTRACTOR"] = args.extractor
    if args.allow_fallback is None:
        args.allow_fallback = bool(args.corpus)

    PERIODS = {} if args.corpus else discover_periods(args.ticker)
    if not args.corpus:
        if not PERIODS:
            print(f"No filings found matching data/raw/*_{args.ticker}.pdf")
            return 2
        args.periods = args.periods or sorted(PERIODS)

    if args.provider:
        os.environ["LLM_PROVIDER"] = args.provider
        os.environ["LLM_FALLBACK_PROVIDER"] = ""   # comparing providers means no silent switch
    if args.api_url:
        # The service decides which model answers; reading LLM_PROVIDER here would tag
        # the scorecard with this process's configuration and quietly mislabel the run.
        import httpx
        try:
            health = httpx.get(args.api_url.rstrip("/") + "/health", timeout=30).json()
        except Exception as exc:
            print(f"cannot reach {args.api_url}: {exc}")
            return 2
        print(f"target: {args.api_url}  ({health['provider']} [{health['input_mode']}], "
              f"commit {health.get('git_commit')})")

        class _Remote:
            name = health["provider"]
            input_mode = health["input_mode"]
        _p = _Remote()
    else:
        # A scorecard is labelled with ONE model, so it must contain only that model's
        # answers. With a fallback chain, a throttled run would otherwise hand some
        # documents to another Gemini model or to Ollama and still be filed under the
        # primary's name -- a score that measured two systems and names one.
        if not args.allow_fallback:
            for key in ("GEMINI_FALLBACK_MODELS", "LLM_FALLBACK_PROVIDER"):
                if os.environ.get(key):
                    print(f"  {key} ignored for evaluation: a scorecard measures one model")
                    os.environ[key] = ""
        else:
            # A corpus pass is long enough that the primary model WILL be throttled, and
            # refusing to fall back would measure the rate limiter rather than the
            # extractor. Every cell records the model that answered it, so the mixture
            # is visible as a slice instead of hidden behind the primary's name.
            print("  fallback chain kept; each cell records which model answered")
        try:
            _p = get_provider_with_fallback()[0]
        except LLMUnavailable as exc:
            # No usable provider -- almost always a reviewer who has not set up a key
            # yet. The committed predictions can still be scored, and that path calls
            # no model at all, so refusing to start would deny them the one thing they
            # can check for free. The tag is rebuilt from configuration instead, which
            # is what names the cache directory anyway.
            provider = os.getenv("LLM_PROVIDER", "gemini").lower()
            model = os.getenv(f"{provider.upper()}_MODEL",
                              "gemini-3.1-flash-lite" if provider == "gemini"
                              else "qwen2.5:7b-instruct")
            mode = os.getenv(f"{provider.upper()}_INPUT_MODE",
                             "image" if provider == "gemini" else "text")

            class _Offline:
                name = f"{provider}:{model}"
                input_mode = mode
            _p = _Offline()

            print(f"⚠ {exc}")
            print(f"  no provider configured, so only CACHED predictions can be scored.")
            print(f"  assuming {_p.name} [{mode}] from configuration, to find the cache.")
            if args.no_cache:
                print("  --no-cache needs a working provider; nothing to extract with.")
                return 2
    provider_tag = provider_tag_for(_p, api_url=args.api_url, extractor=args.extractor)
    # A subset run gets its own filename. Otherwise `--periods Q1` overwrites the
    # full four-period scorecard with a one-period one, and the committed artifact
    # quietly stops meaning what its name says -- which happened once already.
    subset = "" if (args.corpus or set(args.periods) == set(PERIODS)) \
        else "_" + "".join(sorted(args.periods))
    # The ticker names the scorecard too. It was missing for the same reason it was
    # missing from the cache -- one issuer was the only issuer -- and the second one
    # silently overwrote the first one's committed result on its very first run.
    if args.corpus:
        # A subset run gets its own filename, for the reason the single-ticker path
        # already carries one: `--tickers LSIP TAPG` otherwise overwrites the full
        # corpus scorecard with a thirty-five filing one, and the committed artifact
        # quietly stops meaning what its name says. (Done once, in this repo, to the
        # 223-filing scorecard.)
        marks = "".join(sorted(t.upper() for t in (args.tickers or [])))
        if args.limit:
            marks += f"_n{args.limit}s{args.sample_seed}"
        scorecard_name = f"corpus_scorecard_{provider_tag}{'_' + marks if marks else ''}.json"
    else:
        scorecard_name = f"scorecard_{args.ticker}_{provider_tag}{subset}.json"
    print(f"provider: {provider_tag}")

    if args.corpus:
        return run_corpus(args, provider_tag, ROOT / args.out, scorecard_name)

    xlsx = ROOT / "data" / "ground_truth" / f"{args.ticker}.xlsx"
    out_dir = ROOT / args.out
    report = {"ticker": args.ticker, "tolerance": args.tolerance,
              "run": run_metadata(args.tolerance), "periods": {}}
    tally = {"correct": 0, "wrong": 0, "missed": 0, "no_truth": 0, "both_absent": 0}
    failure_kinds = {}
    all_cells, filings, trajectories = [], [], []

    for key in args.periods:
        if key not in PERIODS:
            print(f"[{key}] SKIP - no filing for this period")
            continue
        label, stem, expected_date = PERIODS[key]
        pdf = ROOT / "data" / "raw" / f"{stem}.pdf"
        if not pdf.exists():
            print(f"[{key}] SKIP - {pdf.name} not found")
            continue

        print(f"\n{'='*78}\n{key}  ({label})  {pdf.name}\n{'='*78}")
        truth = load_ground_truth(xlsx, label, ticker=args.ticker)
        ref = PeriodRef(ticker=args.ticker, quarter=key,
                        year=int(expected_date[:4]), pdf=pdf)
        cache_path = cache_path_for(ref, out_dir, provider_tag)

        # One period failing must not abandon the others: a scorecard covering three
        # of four periods with the fourth recorded as failed is still usable, and the
        # recorded failure is itself a result worth keeping.
        try:
            pred, cached, payload = predict(
                pdf, cache_path, not args.no_cache, api_url=args.api_url,
                # The old cache key carried no year, so ARCI Q1 2022 and ARCI Q1 2023
                # were the same file. Entries written under it are still read, but only
                # when the payload's own period_end_date proves which period it holds.
                legacy_path=legacy_cache_path(ref, out_dir, provider_tag),
                extractor=args.extractor)
        except LLMError as exc:
            print(f"  extraction failed: {exc}")
            report["periods"][key] = {"error": str(exc)}
            tally["failed_periods"] = tally.get("failed_periods", 0) + 1
            continue
        if pred is None:
            print("  extraction failed - no fields returned")
            report["periods"][key] = {"error": "extraction_failed"}
            continue
        if cached:
            print("  (cached prediction; --no-cache to re-extract)")

        derived = derived_for(pred, pdf, cache_path, payload)
        rate = derived.fx_rate
        cells = score_period(ref, pred, truth, derived, args.tolerance)
        all_cells += cells
        filings.append({"elapsed_s": getattr(pred, "elapsed_s", None),
                        "usage": getattr(pred, "usage_dict", None)})
        trajectory = getattr(pred, "agent_trajectory", None)
        if trajectory is not None:
            # The filing key travels with the trajectory: page_selection_agreement names
            # the filings whose pages differ, and a list of page numbers with no filing
            # attached cannot be looked at afterwards.
            trajectory.setdefault("key", ref.key)
        trajectories.append(trajectory)

        if pred.period_end_date != expected_date:
            print(f"  \u26a0 period_end_date {pred.period_end_date!r}, expected {expected_date!r}")
        if rate:
            print(f"  kurs {rate:,.2f} IDR/USD")

        print(f"\n  {'field':28}{'as printed':>16}{'truth (USD)':>16}{'IDR err':>10}  status")
        rows = {}
        for cell in cells:
            tally[cell.status] += 1
            if cell.failure_kind:
                failure_kinds[cell.failure_kind] = failure_kinds.get(cell.failure_kind, 0) + 1
            if cell.status == "missed" and cell.abstained:
                tally["missed_abstained"] = tally.get("missed_abstained", 0) + 1

            truth_usd = (cell.truth_idr / rate) if (
                cell.truth_idr and rate and cell.field != "total_share") else cell.truth_idr
            rows[cell.field] = {"status": cell.status, "rel_error": cell.rel_error,
                                "failure_kind": cell.failure_kind,
                                "pred_idr": cell.pred_idr, "truth_idr": cell.truth_idr,
                                "pred_raw": cell.pred_raw, "confidence": cell.confidence,
                                "grounded": cell.grounded, "abstained": cell.abstained}

            mark = {"correct": "OK", "wrong": f"<<< {cell.failure_kind}",
                    "missed": "-- missed", "no_truth": "(no truth)",
                    "both_absent": "(both empty)"}[cell.status]
            if cell.status == "missed" and cell.abstained:
                mark = "-- abstained"
            print(f"  {cell.field:28}"
                  f"{cell.pred_raw if cell.pred_raw is not None else '-':>16}"
                  f"{f'{truth_usd:,.0f}' if truth_usd else '-':>16}"
                  f"{f'{cell.rel_error:.2%}' if cell.rel_error is not None else '-':>10}  {mark}")

        abstained = [f for f, r in rows.items() if r.get("abstained")]
        report["periods"][key] = {"fx_rate": rate, "fields": rows,
                                  "elapsed_s": getattr(pred, "elapsed_s", None),
                                  "usage": getattr(pred, "usage_dict", None),
                                  "abstained": abstained}

    total_scored = tally["correct"] + tally["wrong"] + tally["missed"]
    print(f"\n{'='*78}\nSUMMARY")
    print(f"  correct      {tally['correct']:>3}")
    print(f"  wrong        {tally['wrong']:>3}")
    print(f"  missed       {tally['missed']:>3}   (truth exists, model returned nothing)")
    print(f"  no truth     {tally['no_truth']:>3}   (not scored)")
    if tally.get("missed_abstained"):
        print(f"    of which abstained deliberately: {tally['missed_abstained']}")
    if failure_kinds:
        print("\n  failures by kind:")
        for kind, n in sorted(failure_kinds.items(), key=lambda kv: -kv[1]):
            print(f"    {kind:28}{n:>3}")
    if total_scored:
        print(f"\n  accuracy     {tally['correct']/total_scored:.1%}  ({tally['correct']}/{total_scored})")

    # Totals across the run. Summed from what each period reported rather than
    # recomputed, so a cached period contributes the usage of the run that produced it
    # and a period that failed contributes nothing instead of a zero that reads as free.
    per_period = [p.get("usage") for p in report["periods"].values() if p.get("usage")]
    if per_period:
        tokens_in = [u["tokens"]["input"] for u in per_period
                     if u.get("tokens", {}).get("input") is not None]
        tokens_out = [u["tokens"]["output"] for u in per_period
                      if u.get("tokens", {}).get("output") is not None]
        priced = [u["cost"]["amount"] for u in per_period
                  if u.get("cost", {}).get("amount") is not None]
        totals = {
            "documents": len(per_period),
            "llm_calls": sum(u.get("llm_calls", 0) for u in per_period),
            "tokens": {"input": sum(tokens_in) if tokens_in else None,
                       "output": sum(tokens_out) if tokens_out else None},
            "cost": {"amount": round(sum(priced), 6) if len(priced) == len(per_period)
                                else None,
                     "currency": per_period[0].get("cost", {}).get("currency", "USD"),
                     "reason": None if len(priced) == len(per_period)
                               else per_period[0].get("cost", {}).get("reason")},
        }
        report["usage"] = totals
        print("\n  usage")
        print(f"    documents           {totals['documents']:>3}")
        print(f"    LLM calls           {totals['llm_calls']:>3}")
        t = totals["tokens"]
        print(f"    tokens              {t['input'] if t['input'] is not None else '-'} in"
              f" / {t['output'] if t['output'] is not None else '-'} out")
        cost = totals["cost"]
        if cost["amount"] is not None:
            print(f"    cost                {cost['amount']} {cost['currency']}")
        else:
            print(f"    cost                -  ({cost['reason']})")
        elapsed = [p.get("elapsed_s") for p in report["periods"].values()
                   if p.get("elapsed_s")]
        if elapsed:
            print(f"    latency             {min(elapsed):.0f}-{max(elapsed):.0f}s per "
                  f"document (mean {sum(elapsed)/len(elapsed):.0f}s)")

    report["summary"] = tally
    report["failure_kinds"] = failure_kinds
    report["provider"] = provider_tag
    report["target"] = args.api_url or "in-process"
    # Added beside the old keys, never instead of them: `periods` is what the README
    # documents and what compare_runs has always read, while `cells` is the flat,
    # sliceable form every metric is computed from -- and can be recomputed from, later,
    # by anyone holding the scorecard alone.
    report["schema_version"] = 2
    report["cells"] = [c.to_dict() for c in all_cells]
    report["metrics"] = summarise(all_cells, filings, trajectories)
    report["conversion_sha"] = conversion_sha()
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / scorecard_name).write_text(json.dumps(report, indent=2, default=str))
    print(f"\n  written to {out_dir / scorecard_name}")

    return 1 if tally["wrong"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
